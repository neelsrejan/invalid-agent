"""Week 4 writes a deck whose numbers are the close-out facts."""

from pathlib import Path

from pptx import Presentation

from invalid_agent.graph import main


def test_week_four_writes_a_deck_and_a_projection(tmp_path, monkeypatch):
    monkeypatch.setenv("LLM_BACKEND", "deterministic")
    main(["--scenario", "bad_lot", "--week", "4", "--out", str(tmp_path)])
    closeout = (tmp_path / "closeout.md").read_text(encoding="utf-8")
    assert "Not for clinical use" in closeout
    assert "Not observed follow-up" in closeout
    assert "bad_reagent_lot" in closeout
    deck_path = tmp_path / "exec_deck.pptx"
    assert deck_path.is_file()
    presentation = Presentation(str(deck_path))
    assert len(presentation.slides) == 7
    chunks: list[str] = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                chunks.append(shape.text_frame.text)
    text = "\n".join(chunks)
    assert "Not for clinical use" in text
    assert "Not observed follow-up" in text
    assert "3.67" in text
    assert "L-4419" in text
    assert Path(tmp_path / "closeout_pchart.png").is_file()
