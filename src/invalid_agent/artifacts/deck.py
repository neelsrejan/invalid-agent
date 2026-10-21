"""Seven-slide exec deck. Navy and teal. No model writes the numbers."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

NAVY = RGBColor(0x0B, 0x3A, 0x5B)
TEAL = RGBColor(0x1F, 0x7A, 0x8C)
INK = RGBColor(0x1C, 0x28, 0x30)
PAPER = RGBColor(0xF7, 0xF5, 0xF2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xB2, 0x3A, 0x48)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def _box(slide, left, top, width, height, text, size=18, color=INK, bold=False, fill=None, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(left, top, width, height)
    if fill is not None:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
        shape.line.fill.background()
    frame = shape.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"
    return shape


def _blank(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = PAPER
    _box(slide, 0, 0, SLIDE_W, Inches(0.18), "", fill=NAVY)
    return slide


def _chart_png(facts: dict[str, Any], path: Path) -> Path | None:
    chart = facts.get("chart") or {}
    points = chart.get("points") or []
    if not points:
        return None
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    xs = list(range(len(points)))
    rate = [point["p"] * 100 for point in points]
    ucl = [point["ucl"] * 100 for point in points]
    center = [chart["p_bar"] * 100 for _ in points]
    figure, axis = plt.subplots(figsize=(8.4, 3.2), dpi=120)
    axis.plot(xs, ucl, color="#B23A48", linestyle=":", label="UCL")
    axis.plot(xs, center, color="#8AA0B2", linestyle="--", label="p-bar")
    axis.plot(xs, rate, color="#0B3A5B", marker="o", label="Invalid %")
    axis.set_ylabel("Invalid rate (%)")
    axis.set_title("Observed weekly invalid rate")
    axis.legend(frameon=False, fontsize=8)
    figure.tight_layout()
    path.parent.mkdir(parents=True, exist_ok=True)
    figure.savefig(path)
    plt.close(figure)
    return path


def write_deck(path: Path, facts: dict[str, Any]) -> Path:
    """Write the exec deck. Slide text is the close-out facts, not a new estimate."""
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    prs.core_properties.title = f"Invalid Bot — {facts['scenario_id']}"
    prs.core_properties.subject = "Simulation. Synthetic data. Not for clinical use."

    title = _blank(prs)
    _box(title, Inches(0.6), Inches(1.4), Inches(12), Inches(1.2), "Invalid Bot", size=40, color=NAVY, bold=True)
    _box(
        title,
        Inches(0.6),
        Inches(2.7),
        Inches(12),
        Inches(1.2),
        f"{facts['device']}\nWeek-4 recommendation for {facts['scenario_id']}",
        size=22,
        color=INK,
    )
    _box(
        title,
        Inches(0.6),
        Inches(5.4),
        Inches(12),
        Inches(1.2),
        "Simulation. Synthetic data. Not for clinical use. Not a medical device.",
        size=16,
        color=WHITE,
        bold=True,
        fill=NAVY,
    )

    summary = _blank(prs)
    _box(summary, Inches(0.5), Inches(0.4), Inches(12), Inches(0.6), "Executive summary", size=28, color=NAVY, bold=True)
    _box(
        summary,
        Inches(0.5),
        Inches(1.4),
        Inches(12),
        Inches(4.5),
        (
            f"Observed invalid rate {facts['rate_pct']:.2f}% "
            f"on {facts['k']} of {facts['n']} runs.\n"
            f"Wilson 95% CI {facts['ci_low_pct']:.2f}% to {facts['ci_high_pct']:.2f}%.\n"
            f"Spec is {facts['spec_pct']:.2f}% invalid "
            f"(promised validity {facts['promised_pct']:.2f}%).\n"
            f"Call: {facts['cause']} / {facts['level']} ({facts['category']}).\n"
            f"{facts['recommendation']}"
        ),
        size=22,
    )

    impact = _blank(prs)
    _box(impact, Inches(0.5), Inches(0.4), Inches(12), Inches(0.6), "Impact", size=28, color=NAVY, bold=True)
    _box(
        impact,
        Inches(0.5),
        Inches(1.6),
        Inches(12),
        Inches(3.5),
        (
            f"{facts['k']} invalids of {facts['n']} runs.\n"
            f"${facts['unit']} per invalid, from the scenario file.\n"
            f"Cost of invalids on this file: ${facts['cost']}.\n"
            "That product is computed. It is not a quote from a customer invoice."
        ),
        size=24,
    )

    cause = _blank(prs)
    _box(cause, Inches(0.5), Inches(0.35), Inches(12), Inches(0.5), "Root cause and effectiveness", size=28, color=NAVY, bold=True)
    _box(
        cause,
        Inches(0.5),
        Inches(1.0),
        Inches(12),
        Inches(1.3),
        (
            f"{facts['cause']} / {facts['level']}.\n"
            f"Observed {facts['rate_pct']:.2f}%. Projected {facts['after_pct']:.2f}%.\n"
            f"{facts['label']}"
        ),
        size=18,
    )
    image = _chart_png(facts, path.parent / "closeout_pchart.png")
    if image is not None:
        cause.shapes.add_picture(str(image), Inches(0.7), Inches(2.6), Inches(8.2), Inches(4.2))

    options = _blank(prs)
    _box(options, Inches(0.5), Inches(0.4), Inches(12), Inches(0.6), "Options", size=28, color=NAVY, bold=True)
    contain = (
        f"Contain {facts['level']}. Cost of invalids so far: ${facts['cost']}. "
        "Risk: a second stratum that also cleared the bar would remain in the field."
        if facts["single"]
        else "Do not contain a stratum. No leading stratum cleared the evidence bar."
    )
    _box(options, Inches(0.5), Inches(1.4), Inches(6), Inches(4.5), "Option A\n\n" + contain, size=18, fill=WHITE)
    _box(
        options,
        Inches(6.8),
        Inches(1.4),
        Inches(6),
        Inches(4.5),
        "Option B\n\nWiden the pull before a field action. No second cost is estimated in this file. "
        "Risk: invalids continue while the team waits.",
        size=18,
        fill=WHITE,
    )

    recommend = _blank(prs)
    _box(recommend, Inches(0.5), Inches(0.4), Inches(12), Inches(0.6), "Recommendation", size=28, color=TEAL, bold=True)
    _box(recommend, Inches(0.5), Inches(1.6), Inches(12), Inches(2.2), facts["recommendation"], size=26, color=NAVY, bold=True)
    _box(recommend, Inches(0.5), Inches(4.2), Inches(12), Inches(1.6), facts["d7"], size=18)

    ask = _blank(prs)
    _box(ask, Inches(0.5), Inches(0.4), Inches(12), Inches(0.6), "The ask", size=28, color=NAVY, bold=True)
    _box(
        ask,
        Inches(0.5),
        Inches(1.6),
        Inches(12),
        Inches(2.4),
        "Approve the recommendation this week, or send the case back for the wider pull.",
        size=26,
        color=NAVY,
        bold=True,
    )
    _box(
        ask,
        Inches(0.5),
        Inches(4.4),
        Inches(12),
        Inches(1.6),
        "Do not treat the projected rate as observed follow-up.\n"
        "Simulation. Synthetic data. Not for clinical use. Not a medical device.",
        size=18,
        color=RED,
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return path
